"""構築結果を「メタオリパ 構築」シートと同じ並びの xlsx に書き出す。

PT・合計・還元率などは数式で入れるので、Googleスプレッドシートに取り込んだあと
仕入れ価格や枚数を直しても自動で再計算される。
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .model import PackResult

HEADERS = ["賞", "名称", "URL", "仕入れ価格", "PT構築", "数量", "合計", "配送", "在庫", "当選番号", "確率(%)", "メモ"]
WIDTHS = [8, 34, 40, 12, 12, 9, 14, 10, 6, 22, 9, 30]
YEN = "#,##0"
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
SUMMARY_FILL = PatternFill("solid", fgColor="FFF2CC")
FIRST_ROW = 3


def write_xlsx(result: PackResult, path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "構築"
    _write_build_sheet(ws, result)
    _write_numbers_sheet(wb.create_sheet("当選番号"), result)
    wb.save(path)


def pt_formula(row: int, fee, markup) -> str:
    """構築シートの LET 式と同じ丸めを、LET なしで書いた式（Excel / Googleスプレッドシート両対応）。"""
    x = f"ROUND(D{row}*{fee}*{markup},4)"
    return f"={x}-MOD({x},1000)+IF(MOD({x},1000)<=500,500,1000)"


def _write_build_sheet(ws, result: PackResult) -> None:
    config, lines = result.config, result.lines
    ws["A1"] = config.name
    ws["A1"].font = Font(bold=True, size=13)
    if config.author:
        ws["E1"] = f"作成: {config.author}"

    for col, (header, width) in enumerate(zip(HEADERS, WIDTHS), start=1):
        cell = ws.cell(2, col, header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(col)].width = width

    last = FIRST_ROW + len(lines) - 1
    total_row = last + 1
    for row, line in enumerate(lines, start=FIRST_ROW):
        ws.cell(row, 1, line.tier)
        ws.cell(row, 2, line.name)
        if line.url:
            ws.cell(row, 3, line.url).hyperlink = line.url
        ws.cell(row, 4, line.cost).number_format = YEN
        pt_value = pt_formula(row, config.pricing.fee, line.markup) if line.markup is not None else line.pt
        ws.cell(row, 5, pt_value).number_format = YEN
        ws.cell(row, 6, line.count).number_format = YEN
        ws.cell(row, 7, f"=E{row}*F{row}").number_format = YEN
        ws.cell(row, 8, "配送のみ" if line.ship_only else "")
        ws.cell(row, 9, "在庫" if line.stock else "")
        ws.cell(row, 10, ", ".join(map(str, line.winning_numbers)) if not line.is_filler else "ランダム")
        ws.cell(row, 11, f"=F{row}/$F${total_row}*100").number_format = "0.000"
        ws.cell(row, 12, line.note)

    ws.cell(total_row, 6, f"=SUM(F{FIRST_ROW}:F{last})").number_format = YEN
    ws.cell(total_row, 7, f"=SUM(G{FIRST_ROW}:G{last})").number_format = YEN

    cards_last = FIRST_ROW + len(result.card_lines) - 1

    def rng(col):
        return f"{col}{FIRST_ROW}:{col}{last}"

    def card_rng(col):
        return f"{col}{FIRST_ROW}:{col}{cards_last}"

    r = total_row + 2
    rows = {}
    summary = [
        ("一回", lambda: config.price, YEN),
        ("総口", lambda: f"=F{total_row}", YEN),
        ("総PT", lambda: f"=G{rows['一回']}*G{rows['総口']}", YEN),
        ("構築PT", lambda: f"=G{total_row}", YEN),
        ("還元率", lambda: f"=G{rows['構築PT']}/G{rows['総PT']}", "0.00%"),
        ("差額", lambda: f"=G{rows['総PT']}-G{rows['構築PT']}", YEN),
        ("配送のみ利益", lambda: f'=SUMPRODUCT(({rng("H")}="配送のみ")*({rng("E")}-{rng("D")})*{rng("F")})', YEN),
        ("期待課金額", lambda: f'=SUMIF({rng("H")},"配送のみ",{rng("G")})', YEN),
        ("クーポン利用後課金額", lambda: f"=G{rows['期待課金額']}*{config.pricing.coupon_rate}", YEN),
        (
            "概算収支",
            lambda: f"=G{rows['差額']}+G{rows['配送のみ利益']}+(G{rows['クーポン利用後課金額']}-G{rows['期待課金額']})",
            YEN,
        ),
        (
            "粗利率(決済手数料、納品時手数料込み)",
            lambda: f"=IF(G{rows['クーポン利用後課金額']}=0,0,G{rows['概算収支']}/G{rows['クーポン利用後課金額']})",
            "0.00%",
        ),
        (
            "アド確率",
            lambda: f'=IFERROR("1/"&ROUND(G{rows["総口"]}/SUMIF({rng("E")},">"&G{rows["一回"]},{rng("F")}),1),"なし")',
            None,
        ),
        ("新規仕入額", lambda: f'=SUMPRODUCT(({card_rng("I")}<>"在庫")*{card_rng("D")}*{card_rng("F")})', YEN),
        ("在庫活用額", lambda: f'=SUMPRODUCT(({card_rng("I")}="在庫")*{card_rng("D")}*{card_rng("F")})', YEN),
    ]
    for label, value, fmt in summary:
        rows[label] = r
        ws.cell(r, 6, label).fill = SUMMARY_FILL
        cell = ws.cell(r, 7, value())
        cell.fill = SUMMARY_FILL
        if fmt:
            cell.number_format = fmt
        cell.alignment = Alignment(horizontal="right")
        r += 1

    ws.freeze_panes = "C3"


def _write_numbers_sheet(ws, result: PackResult) -> None:
    """サイト登録用：当選番号順の一覧。"""
    headers = ["当選番号", "賞", "名称", "PT", "配送"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for col, width in enumerate([10, 8, 34, 12, 10], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    entries = sorted(
        (number, line) for line in result.card_lines for number in line.winning_numbers
    )
    for row, (number, line) in enumerate(entries, start=2):
        ws.cell(row, 1, number)
        ws.cell(row, 2, line.tier)
        ws.cell(row, 3, line.name)
        ws.cell(row, 4, line.pt).number_format = YEN
        ws.cell(row, 5, "配送のみ" if line.ship_only else "")
