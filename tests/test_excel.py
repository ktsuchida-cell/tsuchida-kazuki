from pathlib import Path

import openpyxl
import pytest

from oripa import build_pack, load_config, write_xlsx

EXAMPLES = Path(__file__).resolve().parent.parent / "packs" / "examples"


@pytest.fixture
def written(tmp_path):
    result = build_pack(load_config(EXAMPLES / "ad_1in4_1000pt.toml"))
    path = tmp_path / "pack.xlsx"
    write_xlsx(result, path)
    return result, path


def _summary(ws):
    return {ws.cell(r, 6).value: ws.cell(r, 7).value for r in range(1, ws.max_row + 1) if ws.cell(r, 6).value}


def test_layout_matches_build_sheet(written):
    result, path = written
    ws = openpyxl.load_workbook(path)["構築"]
    assert [ws.cell(2, c).value for c in range(1, 8)] == ["賞", "名称", "URL", "仕入れ価格", "PT構築", "数量", "合計"]
    assert ws["A3"].value == "S爆" and ws["H3"].value == "配送のみ"
    assert ws["E3"].value.startswith("=ROUND(D3*1.05*1.4,4)")
    assert ws["E13"].value == 2100  # ハズレ枠は数式ではなく値
    assert {"一回", "総口", "還元率", "配送のみ利益", "概算収支", "アド確率"} <= set(_summary(ws))


def test_numbers_sheet_lists_every_winning_number(written):
    result, path = written
    ws = openpyxl.load_workbook(path)["当選番号"]
    listed = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
    expected = sorted(n for line in result.card_lines for n in line.winning_numbers)
    assert listed == expected


def test_formulas_recalculate_to_builder_metrics(written):
    formulas = pytest.importorskip("formulas")
    result, path = written
    solution = formulas.ExcelModel().loads(str(path)).finish().calculate()
    values = {}
    for key, value in solution.items():
        ref = str(key).split("!")[-1]
        if "構築" in str(key) and ":" not in ref:
            values[ref] = value.value[0][0] if hasattr(value, "value") else value

    ws = openpyxl.load_workbook(path)["構築"]
    rows = {ws.cell(r, 6).value: r for r in range(1, ws.max_row + 1) if ws.cell(r, 6).value}
    m = result.metrics
    for row, line in enumerate(result.lines, start=3):
        assert values[f"E{row}"] == pytest.approx(line.pt)
    assert values[f"G{rows['構築PT']}"] == pytest.approx(m.build_pt)
    assert values[f"G{rows['還元率']}"] == pytest.approx(m.return_rate)
    assert values[f"G{rows['配送のみ利益']}"] == pytest.approx(m.ship_only_profit)
    assert values[f"G{rows['概算収支']}"] == pytest.approx(m.est_profit)
    assert values[f"G{rows['新規仕入額']}"] == pytest.approx(m.new_purchase_cost)
